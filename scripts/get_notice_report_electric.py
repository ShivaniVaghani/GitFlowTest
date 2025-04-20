# Due to DATA availability issue in DEV WHERE t.workflow_id = 219 AND ts.is_deleted = TRUE -> this updated as true

#!/usr/bin/env python3
import argparse
import os
import sys
import base64
from datetime import datetime, timedelta
from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd
from .utils import (
    connect_db,
    metadata_el_nd,
    metadata_el_bd,
    metadata_el_bd_raeco,
    metadata_bd_water,
)

# --- Default date range (previous month window) ---
currect_dir = os.path.dirname(os.path.abspath(__file__))
# Parent directory of the scripts folder
parent_dir = os.path.dirname(currect_dir)
# Default range: first day of previous month to yesterday 23:59:59
yesterday = datetime.now() - timedelta(days=1)
first_day_current_month = yesterday.replace(day=1)
last_day_previous_month = first_day_current_month - timedelta(days=1)
start_prev_month = last_day_previous_month.replace(day=1).strftime("%Y-%m-%d")
end_date_default = yesterday.strftime("%Y-%m-%d, 23:59:59")

# --- Ensure output directory exists (parallel to scripts folder) ---
output_dir = os.path.join(parent_dir, 'output')
os.makedirs(output_dir, exist_ok=True)

# --- Data extraction & transformation functions ---
def get_workflow_metadata_nd_salalah(req_start_date: str = start_prev_month,
                                     req_end_date: str = end_date_default) -> str:
    with connect_db() as conn:
        cursor = conn.cursor()
        query = """
        SELECT DISTINCT
            t.task_id,
            t.account_no,
            t.task_data_json,
            t.task_details_json,
            t.pdf_link,
            u.name,
            t.finished_date,
            u.email,
            ad.account_details_json,
            t.workflow_id,
            ts.is_success,
            ts.outstanding_from_api
        FROM task_completion_info t
        JOIN app_user u ON t.user_id = u.user_id
        JOIN task ts ON t.task_id = ts.task_id
        JOIN account_details ad ON t.account_no = (ad.account_details_json->>'Account No.')::TEXT
        WHERE t.workflow_id = 219 AND ts.is_deleted = TRUE
          AND t.finished_date BETWEEN %s AND %s
        ORDER BY t.finished_date DESC;"""
        cursor.execute(query, (req_start_date, req_end_date))
        results = cursor.fetchall()
        return new_data_to_excel(results)


def get_workflow_metadata_bill_salalah(req_start_date: str = start_prev_month,
                                       req_end_date: str = end_date_default) -> str:
    workflow_ids = [216, 217, 271]
    with connect_db() as conn:
        cursor = conn.cursor()
        query = """
        SELECT DISTINCT
            t.task_id,
            t.account_no,
            t.task_data_json,
            t.task_details_json,
            u.name,
            t.finished_date,
            u.email,
            t.workflow_id
        FROM task_completion_info t
        JOIN app_user u ON t.user_id = u.user_id
        JOIN task ts ON t.task_id = ts.task_id
        WHERE t.workflow_id = ANY(%s)
          AND t.finished_date BETWEEN %s AND %s
        ORDER BY t.finished_date DESC;"""
        cursor.execute(query, (workflow_ids, req_start_date, req_end_date))
        results = cursor.fetchall()
        return bill_data_to_excel(results)

# --- Key renaming & DataFrame builders ---
def update_keys_nd(data: dict, task_details: dict, pdf_link, user_name,
                   master_data, workflowId, successfull, outstanding_from_api) -> dict:
    task_data = {metadata_el_nd.get(k, k): v for k, v in data.items()}
    task_details_map = {
        "User Name": user_name,
        "Activity": workflowId,
        "Successfull": successfull,
        "Remarks": task_details.get("Remarks", ""),
        "Phone": master_data.get("PHONE", " "),
        "Customer Name": master_data.get("CONSUMER_NAME_ENG", " "),
        "Geo Code": task_details.get("Area Code", " "),
        "Updated Outstanding": outstanding_from_api,
        "Consumer Type": task_details.get("Consumer Type", ""),
    }
    return task_data | task_details_map


def update_keys_bd(data: dict, task_details: dict, user_name) -> dict:
    task_data = {metadata_el_bd.get(k, k): v for k, v in data.items()}
    task_details_map = {
        "User Name": user_name,
        "Geo Code": task_details.get("Area Code", " "),
    }
    return task_data | task_details_map


def update_keys_water_bd(data: dict, task_details: dict, user_name) -> dict:
    task_data = {metadata_bd_water.get(k, k): v for k, v in data.items()}
    task_details_map = {
        "User Name": user_name,
        "Geo Code": task_details.get("Area Code", " "),
    }
    return task_data | task_details_map


def update_keys_raeco_bd(data: dict, task_details: dict, user_name) -> dict:
    task_data = {metadata_el_bd_raeco.get(k, k): v for k, v in data.items()}
    task_details_map = {
        "User Name": user_name,
        "Geo Code": task_details.get("Area Code", " "),
    }
    return task_data | task_details_map


def format_hyperlink(url: str) -> str:
    if not pd.isna(url) and url:
        return f'=HYPERLINK("{url}")'
    return ""


def get_unique_nd(df_nd: pd.DataFrame):
    if df_nd.empty:
        return pd.DataFrame({
            "Account Number": [""],
            "Geocode": [""],
            "Activity Intended": [""],
            "Activity Date/Time": [""],
            "Customer Name": [""],
            "Phone": [""],
            "Location": [""],
            "Photo": [""],
            "Field User": [""],
        })
    new_df = pd.DataFrame({
        "Account Number": df_nd.get("accountNo", ""),
        "Geocode": df_nd.get("Geo Code", ""),
        "Activity Intended": "NoticeDelivery",
        "Activity Date/Time": pd.to_datetime(
            df_nd["finishedDate"], dayfirst=False, yearfirst=True
        ).dt.strftime("%Y-%m-%d %H:%M:%S"),
        "Customer Name": df_nd.get("Customer Name", ""),
        "Phone": df_nd.get("Phone", ""),
        "Location": df_nd.get("Location", ""),
        "Photo": df_nd.get("Photo", ""),
        "Field User": df_nd.get("User Name", ""),
    })
    new_df["Photo"] = new_df["Photo"].apply(format_hyperlink)
    return new_df


def get_unique_bd(df_bd: pd.DataFrame):
    if df_bd.empty:
        return pd.DataFrame({
            "Account Number": [""],
            "Field User": [""],
            "Activity Date": [""],
            "Activity Time": [""],
            "GPS Location": [""],
        })
    new_df = pd.DataFrame({
        "Account Number": df_bd.get("accountNo", ""),
        "Field User": df_bd.get("User Name", ""),
        "Activity Date": pd.to_datetime(
            df_bd["finishedDate"], dayfirst=False, yearfirst=True
        ).dt.strftime("%Y-%m-%d"),
        "Activity Time": pd.to_datetime(
            df_bd["finishedDate"], dayfirst=False, yearfirst=True
        ).dt.strftime("%H:%M:%S"),
        "GPS Location": df_bd.get("GPS Location", ""),
    })
    return new_df


def new_data_to_excel(task_data: list[tuple]):
    updated = [
        update_keys_nd(t[2], t[3], t[4], t[5], t[8], t[9], t[10], t[11])
        for t in task_data
    ]
    df = pd.DataFrame(updated)
    uniq = get_unique_nd(df)
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="xlsxwriter") as writer:
        uniq.to_excel(writer, index=False, sheet_name="ND Report")
        ws = writer.sheets["ND Report"]
        table_range = f"A1:{chr(64+len(uniq.columns))}{len(uniq)+1}"
        ws.add_table(table_range, {
            "columns": [{"header": c} for c in uniq.columns],
            "style": "Table Style Medium 16",
            "banded_columns": True
        })
        for i, col in enumerate(uniq.columns):
            width = max(uniq[col].astype(str).map(len).max(), len(col)) + 2
            ws.set_column(i, i, width)
    stream.seek(0)
    return base64.b64encode(stream.read()).decode("utf-8")


def bill_data_to_excel(task_data: list[tuple]):
    bd = [update_keys_bd(t[2], t[3], t[4]) for t in task_data if t[7]==216]
    w  = [update_keys_water_bd(t[2], t[3], t[4]) for t in task_data if t[7]==217]
    r  = [update_keys_raeco_bd(t[2], t[3], t[4]) for t in task_data if t[7]==271]
    dfs = [pd.DataFrame(df) for df in (bd, w, r)]
    names = ["Elec BD Report","Water BD Report","RAECO Elec BD Report"]
    stream = BytesIO()
    with pd.ExcelWriter(stream, engine="xlsxwriter") as writer:
        for name, df in zip(names, dfs):
            uniq = get_unique_bd(df)
            uniq.to_excel(writer, index=False, sheet_name=name)
            ws = writer.sheets[name]
            tbl = f"A1:{chr(64+len(uniq.columns))}{len(uniq)+1}"
            ws.add_table(tbl, {
                "columns":[{"header":c} for c in uniq.columns],
                "style":"Table Style Medium 16",
                "banded_columns":True
            })
            for i, col in enumerate(uniq.columns):
                wdh = max(uniq[col].astype(str).map(len).max(), len(col)) + 2
                ws.set_column(i, i, wdh)
    stream.seek(0)
    return base64.b64encode(stream.read()).decode("utf-8")

# --- CLI entrypoint ---
def main():
    parser = argparse.ArgumentParser(
        description="Generate ND or Bill Excel report and save to output folder"
    )
    parser.add_argument(
        "--type", choices=["nd", "bill"], required=True,
        help="'nd' for Notice Delivery, 'bill' for Bill Delivery"
    )
    parser.add_argument(
        "--startDate", default=start_prev_month,
        help="Start date (YYYY-MM-DD)"
    )
    parser.add_argument(
        "--endDate", default=end_date_default,
        help="End date (YYYY-MM-DD, HH:MM:SS)"
    )

    args = parser.parse_args()

    # Generate Base64 Excel
    if args.type == 'nd':
        b64 = get_workflow_metadata_nd_salalah(args.startDate, args.endDate)
    else:
        b64 = get_workflow_metadata_bill_salalah(args.startDate, args.endDate)

    # Decode and write the Excel file
    excel_data = base64.b64decode(b64)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{args.type}_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)
    with open(filepath, 'wb') as f:
        f.write(excel_data)

    # Print the file path for downstream consumption
    print(filepath)

if __name__ == '__main__':
    main()

